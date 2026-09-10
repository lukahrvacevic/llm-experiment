from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


FULL_SUMMARY = "runs/matrix100-pass5-full-summary.json"
MEDIUM_SUMMARY = "runs/matrix100-pass5-medium-summary.json"
SMALL_GLOB = "runs/matrix100-pass5-small-*-small_context-n100"
DEFAULT_OUTPUT = "runs/repoexec_100tasks_pass1_pass5.xlsx"

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill(fill_type="solid", fgColor="EAF4E2")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export RepoExec matrix summaries to an Excel workbook.")
    parser.add_argument("--full-summary", default=FULL_SUMMARY)
    parser.add_argument("--medium-summary", default=MEDIUM_SUMMARY)
    parser.add_argument("--small-glob", default=SMALL_GLOB)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    return parser.parse_args()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def compute_token_usage(run_dir: Path) -> dict[str, float | int | None]:
    task_metrics_path = run_dir / "task_metrics.jsonl"
    rows = [
        json.loads(line)
        for line in task_metrics_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]

    input_tokens = [row["input_tokens"] for row in rows if row.get("input_tokens") is not None]
    output_tokens = [row["output_tokens"] for row in rows if row.get("output_tokens") is not None]

    return {
        "num_predictions": len(rows),
        "mean_input_tokens_all_predictions": (sum(input_tokens) / len(input_tokens)) if input_tokens else None,
        "mean_output_tokens_all_predictions": (sum(output_tokens) / len(output_tokens)) if output_tokens else None,
        "total_input_tokens_all_predictions": sum(input_tokens) if input_tokens else None,
        "total_output_tokens_all_predictions": sum(output_tokens) if output_tokens else None,
    }


def load_summary_rows(summary_path: Path) -> list[dict[str, Any]]:
    rows = read_json(summary_path)
    for row in rows:
        row["token_usage"] = compute_token_usage(Path(row["output_dir"]))
    return rows


def load_small_rows(glob_pattern: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for run_dir in sorted(Path().glob(glob_pattern)):
        summary = read_json(run_dir / "summary.json")
        timing = read_json(run_dir / "timing.json")
        rows.append(
            {
                "model": timing["model"],
                "subset": timing["subset"],
                "task_limit": timing["task_limit"],
                "output_dir": str(run_dir.resolve()),
                "timing": timing,
                "all_tasks": summary["all_tasks"],
                "cross_context_true": summary["cross_context_true"],
                "token_usage": compute_token_usage(run_dir),
            }
        )
    return rows


def build_context_table(rows: list[dict[str, Any]]) -> list[list[Any]]:
    table: list[list[Any]] = [
        [
            "Model",
            "Pass@1",
            "Pass@5",
            "First Pass@1",
            "DIR",
            "Mean Input Tokens",
            "Mean Output Tokens",
            "Mean Generation Seconds",
            "Peak VRAM MB",
            "Total Seconds",
            "Cross Pass@1",
            "Cross Pass@5",
            "Cross DIR",
        ]
    ]
    for row in sorted(rows, key=lambda item: item["model"]):
        all_tasks = row["all_tasks"]
        cross = row["cross_context_true"]
        timing = row["timing"]
        table.append(
            [
                row["model"],
                all_tasks.get("pass@1"),
                all_tasks.get("pass@5"),
                all_tasks.get("first_pass@1"),
                all_tasks.get("DIR"),
                all_tasks.get("mean_input_tokens"),
                all_tasks.get("mean_output_tokens"),
                all_tasks.get("mean_generation_seconds"),
                all_tasks.get("mean_peak_vram_mb"),
                timing.get("total_seconds"),
                cross.get("pass@1"),
                cross.get("pass@5"),
                cross.get("DIR"),
            ]
        )
    return table


def build_overview_table(
    full_rows: list[dict[str, Any]],
    medium_rows: list[dict[str, Any]],
    small_rows: list[dict[str, Any]],
) -> list[list[Any]]:
    by_context = {
        "full": {row["model"]: row for row in full_rows},
        "medium": {row["model"]: row for row in medium_rows},
        "small": {row["model"]: row for row in small_rows},
    }
    models = sorted(set(by_context["full"]) | set(by_context["medium"]) | set(by_context["small"]))
    table: list[list[Any]] = [
        [
            "Model",
            "Full Pass@1",
            "Full Pass@5",
            "Full DIR",
            "Full Total Seconds",
            "Medium Pass@1",
            "Medium Pass@5",
            "Medium DIR",
            "Medium Total Seconds",
            "Small Pass@1",
            "Small Pass@5",
            "Small DIR",
            "Small Total Seconds",
        ]
    ]

    for model in models:
        row: list[Any] = [model]
        for context in ("full", "medium", "small"):
            item = by_context[context].get(model)
            if item is None:
                row.extend([None, None, None, None])
            else:
                row.extend(
                    [
                        item["all_tasks"].get("pass@1"),
                        item["all_tasks"].get("pass@5"),
                        item["all_tasks"].get("DIR"),
                        item["timing"].get("total_seconds"),
                    ]
                )
        table.append(row)
    return table


def build_token_usage_table(
    full_rows: list[dict[str, Any]],
    medium_rows: list[dict[str, Any]],
    small_rows: list[dict[str, Any]],
) -> list[list[Any]]:
    rows_by_context = {
        "full": {row["model"]: row for row in full_rows},
        "medium": {row["model"]: row for row in medium_rows},
        "small": {row["model"]: row for row in small_rows},
    }
    models = sorted(set(rows_by_context["full"]) | set(rows_by_context["medium"]) | set(rows_by_context["small"]))
    table: list[list[Any]] = [
        [
            "Model",
            "Full Mean In",
            "Full Mean Out",
            "Full Total In",
            "Full Total Out",
            "Medium Mean In",
            "Medium Mean Out",
            "Medium Total In",
            "Medium Total Out",
            "Small Mean In",
            "Small Mean Out",
            "Small Total In",
            "Small Total Out",
        ]
    ]
    for model in models:
        row: list[Any] = [model]
        for context in ("full", "medium", "small"):
            item = rows_by_context[context].get(model)
            if item is None:
                row.extend([None, None, None, None])
                continue
            token_usage = item["token_usage"]
            row.extend(
                [
                    token_usage.get("mean_input_tokens_all_predictions"),
                    token_usage.get("mean_output_tokens_all_predictions"),
                    token_usage.get("total_input_tokens_all_predictions"),
                    token_usage.get("total_output_tokens_all_predictions"),
                ]
            )
        table.append(row)
    return table


def write_table(ws, rows: list[list[Any]], title: str | None = None) -> None:
    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title)
        ws["A1"].font = Font(bold=True, size=13)
        ws["A1"].fill = SECTION_FILL
        start_row = 3

    for row_offset, row_values in enumerate(rows, start=start_row):
        for col_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_offset, column=col_index, value=value)
            if row_offset == start_row:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL

    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(rows[0]))}{start_row + len(rows) - 1}"

    for col_index in range(1, len(rows[0]) + 1):
        max_len = 0
        for row_index in range(start_row, start_row + len(rows)):
            value = ws.cell(row=row_index, column=col_index).value
            text = "" if value is None else str(value)
            max_len = max(max_len, len(text))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max_len + 2, 28)


def apply_number_formats(ws) -> None:
    percentage_keywords = {"pass@", "DIR"}
    seconds_keywords = {"Seconds", "time"}
    for row in ws.iter_rows():
        for cell in row:
            if cell.row < 3:
                continue
            header = ws.cell(row=3 if ws["A1"].value else 1, column=cell.column).value
            if not isinstance(cell.value, (int, float)) or header is None:
                continue
            header_text = str(header)
            if any(keyword in header_text for keyword in percentage_keywords):
                cell.number_format = "0.000"
            elif any(keyword in header_text for keyword in seconds_keywords):
                cell.number_format = "0.00"
            else:
                cell.number_format = "0.00"


def build_workbook(
    full_rows: list[dict[str, Any]],
    medium_rows: list[dict[str, Any]],
    small_rows: list[dict[str, Any]],
) -> Workbook:
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    sheets = [
        ("Overview", build_overview_table(full_rows, medium_rows, small_rows), "RepoExec 100 tasks | Pass@1 and Pass@5"),
        ("Full", build_context_table(full_rows), "Full Context"),
        ("Medium", build_context_table(medium_rows), "Medium Context"),
        ("Small", build_context_table(small_rows), "Small Context"),
        ("Token Usage", build_token_usage_table(full_rows, medium_rows, small_rows), "Token Usage Across Contexts"),
    ]

    for name, rows, title in sheets:
        ws = wb.create_sheet(title=name)
        write_table(ws, rows, title=title)
        apply_number_formats(ws)

    return wb


def main() -> None:
    args = parse_args()
    full_rows = load_summary_rows(Path(args.full_summary))
    medium_rows = load_summary_rows(Path(args.medium_summary))
    small_rows = load_small_rows(args.small_glob)

    workbook = build_workbook(full_rows, medium_rows, small_rows)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(output_path.resolve())


if __name__ == "__main__":
    main()
