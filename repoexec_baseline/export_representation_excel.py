from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


RAW_SUMMARY = "runs/matrix100-pass5-full-summary.json"
AST_GLOB = "runs/matrix100-pass5-full-ast-q4-*-full_context-ast-n100"
REDUCED_GLOB = "runs/matrix100-pass5-full-reduced-ast-q4-*-full_context-reduced_ast-n100"
DEFAULT_OUTPUT = "runs/repoexec_100tasks_raw_ast_reduced.xlsx"

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
SECTION_FILL = PatternFill(fill_type="solid", fgColor="EAF4E2")

MODEL_ORDER = [
    "qwen2.5-coder:1.5b-base",
    "qwen2.5-coder:3b-base",
    "qwen2.5-coder:7b-base",
    "deepseek-coder:1.3b-base-q4_K_M",
    "deepseek-coder:6.7b-base-q4_K_M",
    "deepseek-coder-v2:16b-lite-base-q4_K_M",
    "codegemma:2b-code-q4_K_M",
    "codegemma:7b-code-q4_K_M",
]

RAW_BASELINE_MAP = {
    "deepseek-coder:1.3b-base-q4_K_M": "deepseek-coder:1.3b-base",
    "deepseek-coder:6.7b-base-q4_K_M": "deepseek-coder:6.7b-base",
    "codegemma:2b-code-q4_K_M": "codegemma:2b-code",
    "codegemma:7b-code-q4_K_M": "codegemma:7b-code",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export RepoExec raw vs AST vs reduced AST results to Excel.")
    parser.add_argument("--raw-summary", default=RAW_SUMMARY)
    parser.add_argument("--ast-glob", default=AST_GLOB)
    parser.add_argument("--reduced-glob", default=REDUCED_GLOB)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    return parser.parse_args()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def compute_token_usage(run_dir: Path) -> dict[str, float | int | None]:
    task_metrics_path = run_dir / "task_metrics.jsonl"
    rows = [
        json.loads(line)
        for line in task_metrics_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]

    input_tokens = [row["input_tokens"] for row in rows if row.get("input_tokens") is not None]
    output_tokens = [row["output_tokens"] for row in rows if row.get("output_tokens") is not None]

    return {
        "num_predictions": len(rows),
        "mean_input_tokens_all_predictions": (sum(input_tokens) / len(input_tokens)) if input_tokens else None,
        "mean_output_tokens_all_predictions": (sum(output_tokens) / len(output_tokens)) if output_tokens else None,
        "total_input_tokens_all_predictions": sum(input_tokens) if input_tokens else None,
        "total_output_tokens_all_predictions": sum(output_tokens) if output_tokens else None,
    }


def load_rows(summary_path: Path) -> list[dict[str, Any]]:
    rows = read_json(summary_path)
    for row in rows:
        row["token_usage"] = compute_token_usage(Path(row["output_dir"]))
    return rows


def load_rows_from_glob(glob_pattern: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for run_dir in sorted(Path().glob(glob_pattern)):
        summary = read_json(run_dir / "summary.json")
        timing = read_json(run_dir / "timing.json")
        run_config = read_json(run_dir / "run_config.json")
        rows.append(
            {
                "model": run_config["model"],
                "subset": timing["subset"],
                "representation": run_config.get("representation"),
                "task_limit": timing["task_limit"],
                "output_dir": str(run_dir.resolve()),
                "timing": timing,
                "all_tasks": summary["all_tasks"],
                "cross_context_true": summary["cross_context_true"],
                "token_usage": compute_token_usage(run_dir),
            }
        )
    return rows


def by_model(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {row["model"]: row for row in rows}


def resolve_raw_row(model: str, raw_rows_by_model: dict[str, dict[str, Any]]) -> tuple[dict[str, Any] | None, str | None, bool]:
    if model in raw_rows_by_model:
        return raw_rows_by_model[model], model, True

    mapped = RAW_BASELINE_MAP.get(model)
    if mapped is None:
        return None, None, False
    return raw_rows_by_model.get(mapped), mapped, False


def metric_value(row: dict[str, Any] | None, section: str, key: str) -> Any:
    if row is None:
        return None
    return row.get(section, {}).get(key)


def timing_value(row: dict[str, Any] | None, key: str) -> Any:
    if row is None:
        return None
    return row.get("timing", {}).get(key)


def token_value(row: dict[str, Any] | None, key: str) -> Any:
    if row is None:
        return None
    return row.get("token_usage", {}).get(key)


def build_overview_table(
    raw_rows_by_model: dict[str, dict[str, Any]],
    ast_rows_by_model: dict[str, dict[str, Any]],
    reduced_rows_by_model: dict[str, dict[str, Any]],
) -> list[list[Any]]:
    rows: list[list[Any]] = [[
        "Model",
        "Raw Baseline Model",
        "Raw Exact Match",
        "Raw Pass@1",
        "AST Pass@1",
        "Reduced Pass@1",
        "Raw Pass@5",
        "AST Pass@5",
        "Reduced Pass@5",
        "Raw DIR",
        "AST DIR",
        "Reduced DIR",
        "Raw Total Seconds",
        "AST Total Seconds",
        "Reduced Total Seconds",
    ]]

    for model in MODEL_ORDER:
        raw_row, raw_source, exact_match = resolve_raw_row(model, raw_rows_by_model)
        ast_row = ast_rows_by_model[model]
        reduced_row = reduced_rows_by_model[model]
        rows.append([
            model,
            raw_source,
            "Yes" if exact_match else "No",
            metric_value(raw_row, "all_tasks", "pass@1"),
            metric_value(ast_row, "all_tasks", "pass@1"),
            metric_value(reduced_row, "all_tasks", "pass@1"),
            metric_value(raw_row, "all_tasks", "pass@5"),
            metric_value(ast_row, "all_tasks", "pass@5"),
            metric_value(reduced_row, "all_tasks", "pass@5"),
            metric_value(raw_row, "all_tasks", "DIR"),
            metric_value(ast_row, "all_tasks", "DIR"),
            metric_value(reduced_row, "all_tasks", "DIR"),
            timing_value(raw_row, "total_seconds"),
            timing_value(ast_row, "total_seconds"),
            timing_value(reduced_row, "total_seconds"),
        ])
    return rows


def build_cross_context_table(
    raw_rows_by_model: dict[str, dict[str, Any]],
    ast_rows_by_model: dict[str, dict[str, Any]],
    reduced_rows_by_model: dict[str, dict[str, Any]],
) -> list[list[Any]]:
    rows: list[list[Any]] = [[
        "Model",
        "Raw Baseline Model",
        "Raw Exact Match",
        "Raw Cross Pass@1",
        "AST Cross Pass@1",
        "Reduced Cross Pass@1",
        "Raw Cross Pass@5",
        "AST Cross Pass@5",
        "Reduced Cross Pass@5",
        "Raw Cross DIR",
        "AST Cross DIR",
        "Reduced Cross DIR",
        "Cross Tasks",
    ]]

    for model in MODEL_ORDER:
        raw_row, raw_source, exact_match = resolve_raw_row(model, raw_rows_by_model)
        ast_row = ast_rows_by_model[model]
        reduced_row = reduced_rows_by_model[model]
        rows.append([
            model,
            raw_source,
            "Yes" if exact_match else "No",
            metric_value(raw_row, "cross_context_true", "pass@1"),
            metric_value(ast_row, "cross_context_true", "pass@1"),
            metric_value(reduced_row, "cross_context_true", "pass@1"),
            metric_value(raw_row, "cross_context_true", "pass@5"),
            metric_value(ast_row, "cross_context_true", "pass@5"),
            metric_value(reduced_row, "cross_context_true", "pass@5"),
            metric_value(raw_row, "cross_context_true", "DIR"),
            metric_value(ast_row, "cross_context_true", "DIR"),
            metric_value(reduced_row, "cross_context_true", "DIR"),
            metric_value(ast_row, "cross_context_true", "num_tasks"),
        ])
    return rows


def build_token_usage_table(
    raw_rows_by_model: dict[str, dict[str, Any]],
    ast_rows_by_model: dict[str, dict[str, Any]],
    reduced_rows_by_model: dict[str, dict[str, Any]],
) -> list[list[Any]]:
    rows: list[list[Any]] = [[
        "Model",
        "Raw Baseline Model",
        "Raw Exact Match",
        "Raw Mean Input (summary)",
        "AST Mean Input (summary)",
        "Reduced Mean Input (summary)",
        "Raw Mean Output (summary)",
        "AST Mean Output (summary)",
        "Reduced Mean Output (summary)",
        "Raw Mean Input (all preds)",
        "AST Mean Input (all preds)",
        "Reduced Mean Input (all preds)",
        "Raw Mean Output (all preds)",
        "AST Mean Output (all preds)",
        "Reduced Mean Output (all preds)",
        "Raw Total Input (all preds)",
        "AST Total Input (all preds)",
        "Reduced Total Input (all preds)",
        "Raw Total Output (all preds)",
        "AST Total Output (all preds)",
        "Reduced Total Output (all preds)",
    ]]

    for model in MODEL_ORDER:
        raw_row, raw_source, exact_match = resolve_raw_row(model, raw_rows_by_model)
        ast_row = ast_rows_by_model[model]
        reduced_row = reduced_rows_by_model[model]
        rows.append([
            model,
            raw_source,
            "Yes" if exact_match else "No",
            metric_value(raw_row, "all_tasks", "mean_input_tokens"),
            metric_value(ast_row, "all_tasks", "mean_input_tokens"),
            metric_value(reduced_row, "all_tasks", "mean_input_tokens"),
            metric_value(raw_row, "all_tasks", "mean_output_tokens"),
            metric_value(ast_row, "all_tasks", "mean_output_tokens"),
            metric_value(reduced_row, "all_tasks", "mean_output_tokens"),
            token_value(raw_row, "mean_input_tokens_all_predictions"),
            token_value(ast_row, "mean_input_tokens_all_predictions"),
            token_value(reduced_row, "mean_input_tokens_all_predictions"),
            token_value(raw_row, "mean_output_tokens_all_predictions"),
            token_value(ast_row, "mean_output_tokens_all_predictions"),
            token_value(reduced_row, "mean_output_tokens_all_predictions"),
            token_value(raw_row, "total_input_tokens_all_predictions"),
            token_value(ast_row, "total_input_tokens_all_predictions"),
            token_value(reduced_row, "total_input_tokens_all_predictions"),
            token_value(raw_row, "total_output_tokens_all_predictions"),
            token_value(ast_row, "total_output_tokens_all_predictions"),
            token_value(reduced_row, "total_output_tokens_all_predictions"),
        ])
    return rows


def build_notes_table() -> list[list[Any]]:
    return [
        ["Note", "Meaning"],
        ["Raw Exact Match = No", "Raw baseline comes from the non-q4 variant of the same family, so comparison is approximate."],
        ["Mean Input/Output (summary)", "Mean tokens for the first prediction only, copied from summary.json."],
        ["Mean/Total Input/Output (all preds)", "Computed from task_metrics.jsonl across all 5 predictions per task."],
        ["Cross sheet", "Metrics only for tasks where cross_context == True."],
    ]


def write_table(ws, rows: list[list[Any]], title: str | None = None) -> None:
    start_row = 1
    if title:
        ws.cell(row=1, column=1, value=title)
        ws["A1"].font = Font(bold=True, size=13)
        ws["A1"].fill = SECTION_FILL
        start_row = 3

    for row_offset, row_values in enumerate(rows, start=start_row):
        for col_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_offset, column=col_index, value=value)
            if row_offset == start_row:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL

    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(rows[0]))}{start_row + len(rows) - 1}"

    for col_index in range(1, len(rows[0]) + 1):
        max_len = 0
        for row_index in range(start_row, start_row + len(rows)):
            value = ws.cell(row=row_index, column=col_index).value
            text = "" if value is None else str(value)
            max_len = max(max_len, len(text))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max_len + 2, 28)


def apply_number_formats(ws) -> None:
    percentage_keywords = {"pass@", "DIR"}
    seconds_keywords = {"Seconds", "time"}
    for row in ws.iter_rows():
        for cell in row:
            if cell.row < 3:
                continue
            header = ws.cell(row=3 if ws["A1"].value else 1, column=cell.column).value
            if not isinstance(cell.value, (int, float)) or header is None:
                continue
            header_text = str(header)
            if any(keyword in header_text for keyword in percentage_keywords):
                cell.number_format = "0.000"
            elif any(keyword in header_text for keyword in seconds_keywords):
                cell.number_format = "0.00"
            else:
                cell.number_format = "0.00"


def build_workbook(
    raw_rows: list[dict[str, Any]],
    ast_rows: list[dict[str, Any]],
    reduced_rows: list[dict[str, Any]],
) -> Workbook:
    raw_by_model = by_model(raw_rows)
    ast_by_model = by_model(ast_rows)
    reduced_by_model = by_model(reduced_rows)

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    sheets = [
        ("Overview", build_overview_table(raw_by_model, ast_by_model, reduced_by_model), "Raw vs AST vs Reduced AST"),
        ("Cross Context", build_cross_context_table(raw_by_model, ast_by_model, reduced_by_model), "Cross-Context Comparison"),
        ("Token Usage", build_token_usage_table(raw_by_model, ast_by_model, reduced_by_model), "Token Usage Comparison"),
        ("Notes", build_notes_table(), "Notes"),
    ]

    for name, rows, title in sheets:
        ws = wb.create_sheet(title=name)
        write_table(ws, rows, title=title)
        apply_number_formats(ws)

    return wb


def main() -> None:
    args = parse_args()
    raw_rows = load_rows(Path(args.raw_summary))
    ast_rows = load_rows_from_glob(args.ast_glob)
    reduced_rows = load_rows_from_glob(args.reduced_glob)

    workbook = build_workbook(raw_rows, ast_rows, reduced_rows)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(output_path.resolve())


if __name__ == "__main__":
    main()
